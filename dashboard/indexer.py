import datetime
import io
import random
import streamlit as st


def _tier_info(tier_choice):
    if "Tier 1" in tier_choice:
        return "Tier 1 (Authoritative)", 96
    if "Tier 2" in tier_choice:
        return "Tier 2 (Internal Verified)", 88
    return "Tier 3 (Informal / Draft)", 55


def extract_text(f):
    """Return (text, error) for an uploaded file."""
    name = f.name.lower()
    data = f.getvalue()
    try:
        if name.endswith(".pdf"):
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            return "\n".join((p.extract_text() or "") for p in reader.pages).strip(), ""
        if name.endswith(".docx"):
            import docx
            d = docx.Document(io.BytesIO(data))
            parts = [p.text for p in d.paragraphs]
            for t in d.tables:
                for row in t.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            return "\n".join(parts).strip(), ""
        return data.decode("utf-8", errors="ignore").strip(), ""
    except Exception as e:
        return "", f"{type(e).__name__}: {e}"


def index_files(uploaded_files, tier_choice, dept, chunk_size=512):
    tier_short, rel = _tier_info(tier_choice)
    status = "Certified" if rel >= 75 else "Under Review"
    today = datetime.date.today().strftime("%Y-%m-%d")

    docs = st.session_state.setdefault("documents_list", [])
    raw = st.session_state.setdefault("raw_document_texts", {})
    result = {"new": [], "updated": [], "problems": []}

    for f in uploaded_files:
        text, err = extract_text(f)
        if not text:
            reason = err or "no readable text (scanned/image-only files need OCR)"
            result["problems"].append(f"{f.name}: {reason}")
            continue

        raw[f.name] = text
        size_kb = max(1, round(len(f.getvalue()) / 1024, 1))
        chunks = max(1, len(text) // max(1, chunk_size * 4))  # ~4 chars per token

        existing = next((d for d in docs if d.get("name") == f.name), None)
        if existing:
            existing.update({"chunks": chunks, "size": f"{size_kb} KB", "updated": today})
            result["updated"].append(f.name)
        else:
            docs.insert(0, {
                "id": f"doc-{random.randint(100, 999)}",
                "name": f.name,
                "department": dept,
                "authorityTier": tier_short,
                "reliabilityScore": rel,
                "historicalQueries": 0,
                "positiveFeedback": 0,
                "negativeFeedback": 0,
                "chunks": chunks,
                "size": f"{size_kb} KB",
                "status": "Indexed",
                "trustStatus": status,
                "updated": today,
            })
            result["new"].append(f.name)

    st.session_state["documents_list"] = docs
    st.session_state["total_docs"] = len(docs)
    st.session_state["vector_chunks"] = sum(d.get("chunks", 0) for d in docs)
    return result