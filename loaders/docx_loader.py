from docx import Document
from langchain.schema import Document as LangDocument

def load_docx(file_path):

    doc = Document(file_path)

    text = ""

    for para in doc.paragraphs:
        text += para.text + "\n"

    return [
        LangDocument(
            page_content=text
        )
    ]